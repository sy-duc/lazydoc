"""Test cho các processor cụ thể (TxtProcessor, CsvProcessor, ExcelProcessor, ...)."""

import pytest
from pathlib import Path

from src.processors.txt_processor import TxtProcessor
from src.processors.csv_processor import CsvProcessor
from src.processors.excel_processor import ExcelProcessor
from src.processors.word_processor import WordProcessor
from src.processors.powerpoint_processor import PowerPointProcessor
from src.processors.image_processor import ImageProcessor
from src.processors.factory import ProcessorFactory


class TestTxtProcessor:
    """Test TxtProcessor."""

    def test_extract_utf8(self, tmp_path: Path) -> None:
        f = tmp_path / "test.txt"
        f.write_text("Xin chào thế giới\nDòng thứ hai", encoding="utf-8")
        processor = TxtProcessor()
        result = processor.extract(f)
        assert result.file_name == "test.txt"
        assert "Xin chào thế giới" in result.text_content["main"]
        assert result.metadata["lines"] == 2

    def test_extract_empty_file(self, tmp_path: Path) -> None:
        f = tmp_path / "empty.txt"
        f.write_text("")
        processor = TxtProcessor()
        result = processor.extract(f)
        assert result.text_content["main"] == ""

    def test_supported_extensions(self) -> None:
        processor = TxtProcessor()
        assert processor.supported_extensions == [".txt", ".md"]
        assert processor.can_process(Path("file.txt"))
        assert processor.can_process(Path("file.md"))
        assert not processor.can_process(Path("file.csv"))


class TestCsvProcessor:
    """Test CsvProcessor."""

    def test_extract_basic(self, tmp_path: Path) -> None:
        f = tmp_path / "data.csv"
        f.write_text("name,age\nAlice,30\nBob,25", encoding="utf-8")
        processor = CsvProcessor()
        result = processor.extract(f)
        assert result.metadata["rows"] == 3
        assert result.metadata["columns"] == 2
        assert "main" in result.tables
        table = result.tables["main"][0]
        assert table[0] == ["name", "age"]
        assert table[1] == ["Alice", "30"]

    def test_extract_with_semicolon(self, tmp_path: Path) -> None:
        f = tmp_path / "data.csv"
        f.write_text("name;age\nAlice;30", encoding="utf-8")
        processor = CsvProcessor()
        result = processor.extract(f)
        assert result.metadata["rows"] == 2


class TestExcelProcessor:
    """Test ExcelProcessor."""

    def test_extract_xlsx(self, tmp_path: Path) -> None:
        import openpyxl
        f = tmp_path / "test.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "Tên"
        ws["B1"] = "Tuổi"
        ws["A2"] = "Alice"
        ws["B2"] = 30
        wb.save(f)

        processor = ExcelProcessor()
        result = processor.extract(f)
        assert result.metadata["sheets"] == 1
        # Dữ liệu Excel nằm trong tables (không còn text_content để tránh gửi 2 lần lên AI)
        assert "Sheet1" in result.tables
        sheet_table = result.tables["Sheet1"][0]
        assert any("Alice" in row for row in sheet_table)

    def test_extract_xlsx_multi_sheet(self, tmp_path: Path) -> None:
        import openpyxl
        f = tmp_path / "multi.xlsx"
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Dữ liệu"
        ws1["A1"] = "Hello"
        ws2 = wb.create_sheet("Tổng hợp")
        ws2["A1"] = "Tổng"
        wb.save(f)

        processor = ExcelProcessor()
        result = processor.extract(f)
        assert result.metadata["sheets"] == 2
        assert "Dữ liệu" in result.tables
        assert "Tổng hợp" in result.tables

    def test_extract_xlsx_merge_cells(self, tmp_path: Path) -> None:
        import openpyxl
        f = tmp_path / "merge.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Tiêu đề gộp"
        ws.merge_cells("A1:C1")
        ws["A2"] = "Dữ liệu"
        wb.save(f)

        processor = ExcelProcessor()
        result = processor.extract(f)
        sheet_name = result.metadata["sheet_names"]
        assert sheet_name in result.tables
        table_rows = result.tables[sheet_name][0]
        all_cells = [cell for row in table_rows for cell in row]
        assert "Tiêu đề gộp" in all_cells

    def test_supported_extensions(self) -> None:
        processor = ExcelProcessor()
        assert ".xlsx" in processor.supported_extensions
        assert ".xls" in processor.supported_extensions


class TestWordProcessor:
    """Test WordProcessor."""

    def test_extract_docx(self, tmp_path: Path) -> None:
        import docx
        f = tmp_path / "test.docx"
        doc = docx.Document()
        doc.add_paragraph("Đoạn văn thứ nhất")
        doc.add_paragraph("Đoạn văn thứ hai")
        doc.save(f)

        processor = WordProcessor()
        result = processor.extract(f)
        # Prose text lưu dưới key "prose" (tách riêng với "tables" để tránh trùng lặp)
        assert "Đoạn văn thứ nhất" in result.text_content["prose"]
        assert "Đoạn văn thứ hai" in result.text_content["prose"]

    def test_extract_docx_with_table(self, tmp_path: Path) -> None:
        import docx
        f = tmp_path / "table.docx"
        doc = docx.Document()
        doc.add_paragraph("Tiêu đề")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "2"
        doc.save(f)

        processor = WordProcessor()
        result = processor.extract(f)
        assert result.metadata["tables"] == 1
        assert "tables" in result.tables


class TestPowerPointProcessor:
    """Test PowerPointProcessor."""

    def test_extract_pptx(self, tmp_path: Path) -> None:
        from pptx import Presentation
        f = tmp_path / "test.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Tiêu đề slide"
        slide.placeholders[1].text = "Nội dung slide"
        prs.save(f)

        processor = PowerPointProcessor()
        result = processor.extract(f)
        assert result.metadata["slides"] == 1
        assert "Slide 1" in result.text_content
        assert "Tiêu đề slide" in result.text_content["Slide 1"]

    def test_extract_pptx_multi_slides(self, tmp_path: Path) -> None:
        from pptx import Presentation
        f = tmp_path / "multi.pptx"
        prs = Presentation()
        for i in range(3):
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = f"Slide {i + 1}"
        prs.save(f)

        processor = PowerPointProcessor()
        result = processor.extract(f)
        assert result.metadata["slides"] == 3


class TestImageProcessor:
    """Test ImageProcessor."""

    def test_extract_png(self, tmp_path: Path) -> None:
        from PIL import Image
        f = tmp_path / "test.png"
        img = Image.new("RGB", (100, 50), color="red")
        img.save(f)

        processor = ImageProcessor()
        result = processor.extract(f)
        assert result.metadata["width"] == 100
        assert result.metadata["height"] == 50
        assert f.name in result.images
        assert len(result.images[f.name]) > 0

    def test_extract_jpeg(self, tmp_path: Path) -> None:
        from PIL import Image
        f = tmp_path / "test.jpg"
        img = Image.new("RGB", (200, 100), color="blue")
        img.save(f)

        processor = ImageProcessor()
        result = processor.extract(f)
        assert result.metadata["width"] == 200
        assert result.metadata["height"] == 100

    def test_extract_rgba(self, tmp_path: Path) -> None:
        from PIL import Image
        f = tmp_path / "test.png"
        img = Image.new("RGBA", (50, 50), color=(255, 0, 0, 128))
        img.save(f)

        processor = ImageProcessor()
        result = processor.extract(f)
        assert result.has_content


class TestProcessorFactoryIntegration:
    """Test tích hợp: Factory trả về đúng processor và extract thành công."""

    def test_txt_via_factory(self, tmp_path: Path) -> None:
        f = tmp_path / "test.txt"
        f.write_text("hello world")
        processor = ProcessorFactory.get_processor(f)
        assert isinstance(processor, TxtProcessor)
        result = processor.extract(f)
        assert result.has_content

    def test_csv_via_factory(self, tmp_path: Path) -> None:
        f = tmp_path / "data.csv"
        f.write_text("a,b\n1,2")
        processor = ProcessorFactory.get_processor(f)
        assert isinstance(processor, CsvProcessor)
        result = processor.extract(f)
        assert result.has_content

    def test_xlsx_via_factory(self, tmp_path: Path) -> None:
        import openpyxl
        f = tmp_path / "test.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"] = "test"
        wb.save(f)
        processor = ProcessorFactory.get_processor(f)
        assert isinstance(processor, ExcelProcessor)
        result = processor.extract(f)
        assert result.has_content

    def test_docx_via_factory(self, tmp_path: Path) -> None:
        import docx
        f = tmp_path / "test.docx"
        doc = docx.Document()
        doc.add_paragraph("test")
        doc.save(f)
        processor = ProcessorFactory.get_processor(f)
        assert isinstance(processor, WordProcessor)
        result = processor.extract(f)
        assert result.has_content

    def test_pptx_via_factory(self, tmp_path: Path) -> None:
        from pptx import Presentation
        f = tmp_path / "test.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "test"
        prs.save(f)
        processor = ProcessorFactory.get_processor(f)
        assert isinstance(processor, PowerPointProcessor)
        result = processor.extract(f)
        assert result.has_content

    def test_image_via_factory(self, tmp_path: Path) -> None:
        from PIL import Image
        f = tmp_path / "test.png"
        Image.new("RGB", (10, 10)).save(f)
        processor = ProcessorFactory.get_processor(f)
        assert isinstance(processor, ImageProcessor)
        result = processor.extract(f)
        assert result.has_content
