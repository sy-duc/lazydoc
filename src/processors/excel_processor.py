"""ExcelProcessor — Trích xuất nội dung từ file Excel (.xlsx, .xls)."""

import logging
from pathlib import Path

from src.processors.base import ExtractedContent, FileProcessor

logger = logging.getLogger(__name__)


class ExcelProcessor(FileProcessor):
    """Processor cho file Excel (.xlsx, .xls).

    Hỗ trợ nhiều sheet, merge cells, shapes (XML parsing), hình ảnh.
    """

    @property
    def supported_extensions(self) -> list[str]:
        return [".xlsx", ".xls"]

    def extract(self, file_path: Path) -> ExtractedContent:
        """Trích xuất nội dung từ file Excel.

        Args:
            file_path: Đường dẫn file Excel.

        Returns:
            ExtractedContent với dữ liệu từng sheet.
        """
        self.validate_file(file_path)

        ext = file_path.suffix.lower()
        if ext == ".xls":
            return self._extract_xls(file_path)
        return self._extract_xlsx(file_path)

    def _extract_xlsx(self, file_path: Path) -> ExtractedContent:
        """Trích xuất từ file .xlsx bằng openpyxl + XML parsing cho shapes."""
        import openpyxl
        from src.processors.xml_shapes import (
            extract_shapes_from_xlsx,
            extract_images_from_xlsx,
            map_drawings_to_sheets,
        )

        wb = openpyxl.load_workbook(file_path, data_only=True)

        text_content: dict[str, str] = {}
        tables: dict[str, list[list[list[str]]]] = {}
        shapes_text: dict[str, list[str]] = {}

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # Extract dữ liệu cell (xử lý merge cells)
            rows_data: list[list[str]] = []
            merged_values = self._get_merged_cell_values(ws)

            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
                row_data: list[str] = []
                for cell in row:
                    coord = cell.coordinate
                    if coord in merged_values:
                        row_data.append(str(merged_values[coord]))
                    elif cell.value is not None:
                        row_data.append(str(cell.value))
                    else:
                        row_data.append("")
                rows_data.append(row_data)

            # Lưu vào tables — text_content bỏ qua để tránh gửi cùng dữ liệu 2 lần lên AI
            non_empty_rows = [r for r in rows_data if any(c.strip() for c in r)]
            if non_empty_rows:
                tables[sheet_name] = [rows_data]

        wb.close()

        # Extract shapes text bằng XML parsing (zipfile)
        drawing_to_sheet = map_drawings_to_sheets(file_path)
        shapes_by_drawing = extract_shapes_from_xlsx(file_path)
        for drawing_name, texts in shapes_by_drawing.items():
            sheet_name = drawing_to_sheet.get(drawing_name, drawing_name)
            if sheet_name in shapes_text:
                shapes_text[sheet_name].extend(texts)
            else:
                shapes_text[sheet_name] = texts

        # Extract hình ảnh nhúng từ xl/media/
        images = extract_images_from_xlsx(file_path)

        # Đếm shapes
        total_shapes = sum(len(v) for v in shapes_text.values())

        content = ExtractedContent(
            file_path=file_path,
            file_name=file_path.name,
            file_size=file_path.stat().st_size,
            text_content=text_content,
            tables=tables,
            shapes_text=shapes_text,
            images=images,
            metadata={
                "sheets": len(wb.sheetnames),
                "sheet_names": ", ".join(wb.sheetnames),
                "shapes": total_shapes,
                "images": len(images),
            },
        )

        logger.info(
            "Đã extract file xlsx: %s (%d sheet, %d shapes, %d ảnh)",
            file_path.name, len(wb.sheetnames), total_shapes, len(images),
        )
        return content

    def _get_merged_cell_values(self, ws: "openpyxl.worksheet.worksheet.Worksheet") -> dict[str, str]:
        """Lấy giá trị cho các merged cells.

        Args:
            ws: Worksheet cần xử lý.

        Returns:
            Dict mapping coordinate → giá trị cho mọi cell trong merged range.
        """
        merged_values: dict[str, str] = {}
        for merged_range in ws.merged_cells.ranges:
            top_left_value = ws.cell(merged_range.min_row, merged_range.min_col).value
            value_str = str(top_left_value) if top_left_value is not None else ""
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    cell = ws.cell(row, col)
                    merged_values[cell.coordinate] = value_str
        return merged_values

    def _extract_xls(self, file_path: Path) -> ExtractedContent:
        """Trích xuất từ file .xls bằng xlrd.

        Lưu ý: xlrd không hỗ trợ shapes/images. Chỉ extract được cell data.
        """
        import xlrd

        wb = xlrd.open_workbook(file_path)

        text_content: dict[str, str] = {}
        tables: dict[str, list[list[list[str]]]] = {}

        for sheet_idx in range(wb.nsheets):
            ws = wb.sheet_by_index(sheet_idx)
            sheet_name = ws.name

            rows_data: list[list[str]] = []
            for row_idx in range(ws.nrows):
                row_data: list[str] = []
                for col_idx in range(ws.ncols):
                    cell_value = ws.cell_value(row_idx, col_idx)
                    row_data.append(str(cell_value) if cell_value != "" else "")
                rows_data.append(row_data)

            text_lines = []
            for row in rows_data:
                line = "\t".join(row)
                if line.strip():
                    text_lines.append(line)

            if text_lines:
                text_content[sheet_name] = "\n".join(text_lines)
                tables[sheet_name] = [rows_data]

        content = ExtractedContent(
            file_path=file_path,
            file_name=file_path.name,
            file_size=file_path.stat().st_size,
            text_content=text_content,
            tables=tables,
            metadata={"sheets": wb.nsheets, "sheet_names": ", ".join(wb.sheet_names())},
        )

        logger.info("Đã extract file xls: %s (%d sheet)", file_path.name, wb.nsheets)
        return content
