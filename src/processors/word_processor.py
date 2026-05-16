"""WordProcessor — Trích xuất nội dung từ file Word (.docx, .doc)."""

import logging
from pathlib import Path

from src.processors.base import ExtractedContent, FileProcessor

logger = logging.getLogger(__name__)


class WordProcessor(FileProcessor):
    """Processor cho file Word (.docx, .doc).

    Hỗ trợ bảng, hình ảnh, shapes/textbox (XML parsing).
    File .doc được xử lý nếu có LibreOffice (chuyển đổi → .docx).
    """

    @property
    def supported_extensions(self) -> list[str]:
        return [".docx", ".doc"]

    def extract(self, file_path: Path) -> ExtractedContent:
        """Trích xuất nội dung từ file Word.

        Args:
            file_path: Đường dẫn file Word.

        Returns:
            ExtractedContent với text, bảng, shapes, hình ảnh.
        """
        self.validate_file(file_path)

        ext = file_path.suffix.lower()
        if ext == ".doc":
            return self._extract_doc(file_path)
        return self._extract_docx(file_path)

    def _extract_docx(self, file_path: Path) -> ExtractedContent:
        """Trích xuất từ file .docx bằng python-docx + XML parsing cho shapes."""
        import docx
        from src.processors.xml_shapes import extract_shapes_from_docx

        doc = docx.Document(file_path)

        text_parts: list[str] = []
        tables_data: list[list[list[str]]] = []
        images: dict[str, bytes] = {}

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Extract tables — chỉ lưu vào tables_data, không thêm vào text_parts
        # để tránh gửi cùng dữ liệu 2 lần lên AI (text_content + tables)
        for idx, table in enumerate(doc.tables):
            table_rows: list[list[str]] = []
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                table_rows.append(row_data)
            tables_data.append(table_rows)

        # Extract images
        for rel_id, rel in doc.part.rels.items():
            if "image" in rel.reltype:
                try:
                    image_data = rel.target_part.blob
                    image_name = Path(rel.target_ref).name
                    images[image_name] = image_data
                except Exception as e:
                    logger.warning("Không thể đọc hình ảnh %s: %s", rel_id, e)

        # Extract shapes/textbox text bằng XML parsing
        shapes_text_list = extract_shapes_from_docx(file_path)

        text_content = "\n".join(text_parts)

        content = ExtractedContent(
            file_path=file_path,
            file_name=file_path.name,
            file_size=file_path.stat().st_size,
            text_content={"prose": text_content} if text_content.strip() else {},
            tables={"tables": tables_data} if tables_data else {},
            shapes_text={"main": shapes_text_list} if shapes_text_list else {},
            images=images,
            metadata={
                "paragraphs": len(doc.paragraphs),
                "tables": len(doc.tables),
                "images": len(images),
                "shapes": len(shapes_text_list),
            },
        )

        logger.info(
            "Đã extract file docx: %s (%d đoạn, %d bảng, %d ảnh, %d shapes)",
            file_path.name, len(doc.paragraphs), len(doc.tables),
            len(images), len(shapes_text_list),
        )
        return content

    def _extract_doc(self, file_path: Path) -> ExtractedContent:
        """Trích xuất từ file .doc (chuyển đổi sang .docx trước).

        Yêu cầu LibreOffice được cài đặt trên hệ thống.
        """
        import subprocess
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir:
            try:
                subprocess.run(
                    ["libreoffice", "--headless", "--convert-to", "docx", str(file_path), "--outdir", tmp_dir],
                    capture_output=True,
                    timeout=60,
                    check=True,
                )
            except FileNotFoundError:
                raise RuntimeError(
                    "Cần cài LibreOffice để xử lý file .doc. "
                    "Vui lòng cài đặt hoặc chuyển file sang .docx."
                )
            except subprocess.TimeoutExpired:
                raise RuntimeError(f"Chuyển đổi file .doc quá lâu (>60s): {file_path.name}")

            docx_path = Path(tmp_dir) / file_path.with_suffix(".docx").name
            if not docx_path.exists():
                raise RuntimeError(f"Không thể chuyển đổi file .doc: {file_path.name}")

            result = self._extract_docx(docx_path)
            result.file_path = file_path
            result.file_name = file_path.name
            result.file_size = file_path.stat().st_size
            return result
