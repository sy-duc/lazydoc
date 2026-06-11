"""ExcelWriter — Ghi file Excel đã dịch, giữ nguyên định dạng gốc."""

import logging
import re
import shutil
import zipfile
from pathlib import Path
from typing import Callable
from xml.sax.saxutils import escape as xml_escape

from src.core.logging_config import safe_file_label
from src.writers.base import FileWriter

logger = logging.getLogger(__name__)

# Namespace chính của xlsx
_SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


class ExcelWriter(FileWriter):
    """Writer cho file Excel (.xlsx, .xls).

    Chiến lược 2 nhánh:
    - File có shapes/images: thao tác tầng ZIP, dùng regex sửa XML text
      mà không parse/serialize lại (giữ nguyên namespace, cấu trúc).
    - File đơn giản: dùng openpyxl (nhanh hơn, dễ xử lý formatting).
    """

    @property
    def supported_extensions(self) -> list[str]:
        return [".xlsx", ".xls"]

    def write_translated(
        self,
        source_path: Path,
        output_path: Path,
        translate_fn: Callable[[str], str],
    ) -> None:
        """Dịch file Excel giữ nguyên định dạng.

        Args:
            source_path: Đường dẫn file gốc.
            output_path: Đường dẫn file đầu ra.
            translate_fn: Hàm dịch.
        """
        ext = source_path.suffix.lower()
        if ext == ".xls":
            self._write_xls(source_path, output_path, translate_fn)
        elif self._has_shapes_or_media(source_path):
            self._write_xlsx_zip(source_path, output_path, translate_fn)
        else:
            self._write_xlsx_openpyxl(source_path, output_path, translate_fn)

    # --- Detect ---

    @staticmethod
    def _has_shapes_or_media(source_path: Path) -> bool:
        """Kiểm tra file xlsx có chứa shapes, images, hoặc drawings không."""
        try:
            with zipfile.ZipFile(source_path, "r") as zf:
                for name in zf.namelist():
                    if name.startswith("xl/drawings/") or name.startswith("xl/media/"):
                        return True
        except zipfile.BadZipFile:
            pass
        return False

    # --- Nhánh ZIP + regex (giữ nguyên shapes) ---

    def _write_xlsx_zip(
        self,
        source_path: Path,
        output_path: Path,
        translate_fn: Callable[[str], str],
    ) -> None:
        """Dịch file .xlsx bằng thao tác ZIP + regex — giữ nguyên shapes/images.

        Dùng regex thay thế text trong XML thay vì parse/serialize lại,
        tránh phá vỡ namespace prefixes và cấu trúc XML gốc.
        """
        shutil.copy2(source_path, output_path)

        with zipfile.ZipFile(source_path, "r") as zf_in:
            modified_files: dict[str, bytes] = {}

            # Dịch sharedStrings.xml (chứa toàn bộ text cells)
            if "xl/sharedStrings.xml" in zf_in.namelist():
                sst_xml = zf_in.read("xl/sharedStrings.xml")
                translated_sst = self._translate_shared_strings_regex(sst_xml, translate_fn)
                modified_files["xl/sharedStrings.xml"] = translated_sst

            # Dịch tên sheet trong workbook.xml
            if "xl/workbook.xml" in zf_in.namelist():
                wb_xml = zf_in.read("xl/workbook.xml")
                translated_wb = self._translate_sheet_names_regex(wb_xml, translate_fn)
                modified_files["xl/workbook.xml"] = translated_wb

            # Dịch inline strings trong từng sheet
            for name in zf_in.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    sheet_xml = zf_in.read(name)
                    translated_sheet = self._translate_inline_strings_regex(
                        sheet_xml, translate_fn
                    )
                    if translated_sheet is not None:
                        modified_files[name] = translated_sheet

            # Dịch text trong shapes (drawings), bỏ qua charts
            for name in zf_in.namelist():
                if name.startswith("xl/drawings/") and name.endswith(".xml"):
                    drawing_xml = zf_in.read(name)
                    translated_drawing = self._translate_drawing_text_regex(
                        drawing_xml, translate_fn
                    )
                    if translated_drawing is not None:
                        modified_files[name] = translated_drawing

        self._replace_in_zip(output_path, modified_files)
        logger.info("Đã ghi file xlsx (ZIP mode): %s", safe_file_label(output_path))

    @staticmethod
    def _translate_shared_strings_regex(
        xml_data: bytes, translate_fn: Callable[[str], str]
    ) -> bytes:
        """Dịch text trong sharedStrings.xml bằng regex.

        Tìm tất cả thẻ <t>...</t> (có hoặc không có namespace prefix)
        và dịch nội dung text bên trong, giữ nguyên XML xung quanh.
        """
        xml_str = xml_data.decode("utf-8")

        # Match thẻ <t> với mọi namespace prefix: <t>, <x:t>, <ns0:t>, ...
        # Cũng match attributes như xml:space="preserve"
        pattern = re.compile(r'(<(?:[\w.:]+)?t(?:\s[^>]*)?>)([^<]+)(</(?:[\w.:]+)?t>)')

        def _replace_text(match: re.Match) -> str:
            open_tag = match.group(1)
            text = match.group(2)
            close_tag = match.group(3)
            if text.strip():
                text = xml_escape(translate_fn(text))
            return f"{open_tag}{text}{close_tag}"

        translated = pattern.sub(_replace_text, xml_str)
        return translated.encode("utf-8")

    @staticmethod
    def _translate_sheet_names_regex(
        xml_data: bytes, translate_fn: Callable[[str], str]
    ) -> bytes:
        """Dịch tên sheet trong workbook.xml bằng regex."""
        xml_str = xml_data.decode("utf-8")

        # Match attribute name="..." trong thẻ <sheet ...>
        pattern = re.compile(r'(<(?:[\w.:]+)?sheet\s[^>]*\bname=")([^"]+)(")')

        def _replace_name(match: re.Match) -> str:
            prefix = match.group(1)
            name = match.group(2)
            suffix = match.group(3)
            if name.strip():
                translated = translate_fn(name)
                # Sanitize tên sheet (Excel và XML attribute đều không cho phép " hay \[]*?/\)
                translated = re.sub(r'[\[\]*?/\\":]', '', translated)
                translated = translated[:31].strip() or "Sheet"
                return f"{prefix}{translated}{suffix}"
            return match.group(0)

        translated = pattern.sub(_replace_name, xml_str)
        return translated.encode("utf-8")

    @staticmethod
    def _translate_inline_strings_regex(
        xml_data: bytes, translate_fn: Callable[[str], str]
    ) -> bytes | None:
        """Dịch inline strings trong sheet XML bằng regex.

        Inline strings nằm trong thẻ <is><t>text</t></is>.
        """
        xml_str = xml_data.decode("utf-8")

        # Chỉ match <t> bên trong <is>...</is>
        is_pattern = re.compile(
            r'(<(?:[\w.:]+)?is(?:\s[^>]*)?>)(.*?)(</(?:[\w.:]+)?is>)',
            re.DOTALL,
        )

        has_changes = False

        def _replace_is_block(match: re.Match) -> str:
            nonlocal has_changes
            open_is = match.group(1)
            inner = match.group(2)
            close_is = match.group(3)

            t_pattern = re.compile(r'(<(?:[\w.:]+)?t(?:\s[^>]*)?>)([^<]+)(</(?:[\w.:]+)?t>)')

            def _replace_t(t_match: re.Match) -> str:
                nonlocal has_changes
                open_t = t_match.group(1)
                text = t_match.group(2)
                close_t = t_match.group(3)
                if text.strip():
                    has_changes = True
                    text = xml_escape(translate_fn(text))
                return f"{open_t}{text}{close_t}"

            new_inner = t_pattern.sub(_replace_t, inner)
            return f"{open_is}{new_inner}{close_is}"

        translated = is_pattern.sub(_replace_is_block, xml_str)

        if not has_changes:
            return None
        return translated.encode("utf-8")

    @staticmethod
    def _translate_drawing_text_regex(
        xml_data: bytes, translate_fn: Callable[[str], str]
    ) -> bytes | None:
        """Dịch text trong drawing XML (shapes, textboxes) bằng regex.

        Drawing XML chứa text trong thẻ <a:t>text</a:t> (namespace drawingML).
        Chỉ dịch text, giữ nguyên mọi thứ khác (hình ảnh, biểu đồ ref, ...).
        """
        xml_str = xml_data.decode("utf-8")

        # Match thẻ <a:t>text</a:t> hoặc <*:t>text</*:t> trong drawing
        pattern = re.compile(r'(<(?:[\w.:]+)?t(?:\s[^>]*)?>)([^<]+)(</(?:[\w.:]+)?t>)')

        has_changes = False

        def _replace_text(match: re.Match) -> str:
            nonlocal has_changes
            open_tag = match.group(1)
            text = match.group(2)
            close_tag = match.group(3)
            if text.strip():
                has_changes = True
                text = xml_escape(translate_fn(text))
            return f"{open_tag}{text}{close_tag}"

        translated = pattern.sub(_replace_text, xml_str)

        if not has_changes:
            return None
        return translated.encode("utf-8")

    @staticmethod
    def _replace_in_zip(zip_path: Path, modified_files: dict[str, bytes]) -> None:
        """Thay thế các file trong zip mà không ảnh hưởng file khác."""
        if not modified_files:
            return

        temp_path = zip_path.with_suffix(".tmp")
        with zipfile.ZipFile(zip_path, "r") as zf_in, \
             zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as zf_out:
            for item in zf_in.infolist():
                if item.filename in modified_files:
                    zf_out.writestr(item, modified_files[item.filename])
                else:
                    zf_out.writestr(item, zf_in.read(item.filename))

        temp_path.replace(zip_path)

    # --- Nhánh openpyxl (file đơn giản) ---

    @staticmethod
    def _sanitize_sheet_name(name: str) -> str:
        """Chuẩn hóa tên sheet theo quy tắc Excel."""
        name = re.sub(r'[\[\]*?/\\":]', '', name)
        name = name[:31].strip()
        return name if name else "Sheet"

    def _write_xlsx_openpyxl(
        self,
        source_path: Path,
        output_path: Path,
        translate_fn: Callable[[str], str],
    ) -> None:
        """Dịch file .xlsx bằng openpyxl — cho file không có shapes."""
        import openpyxl

        wb = openpyxl.load_workbook(source_path)

        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None and isinstance(cell.value, str):
                        original = cell.value
                        if original.strip():
                            cell.value = translate_fn(original)

        # Dịch tên sheet
        for ws in wb.worksheets:
            original_title = ws.title
            if original_title.strip():
                translated_title = self._sanitize_sheet_name(translate_fn(original_title))
                ws.title = translated_title

        wb.save(output_path)
        wb.close()
        logger.info("Đã ghi file xlsx (openpyxl mode): %s", safe_file_label(output_path))

    # --- Nhánh xls (legacy) ---

    def _write_xls(
        self,
        source_path: Path,
        output_path: Path,
        translate_fn: Callable[[str], str],
    ) -> None:
        """Dịch file .xls — chuyển sang .xlsx vì xlrd chỉ đọc.

        File .xls được đọc bằng xlrd, ghi ra .xlsx bằng openpyxl.
        Giữ được dữ liệu nhưng có thể mất một số formatting phức tạp.
        """
        import xlrd
        import openpyxl

        rb = xlrd.open_workbook(source_path)
        wb = openpyxl.Workbook()

        for sheet_idx in range(rb.nsheets):
            rs = rb.sheet_by_index(sheet_idx)

            if sheet_idx == 0:
                ws = wb.active
            else:
                ws = wb.create_sheet()

            # Dịch tên sheet
            if rs.name.strip():
                ws.title = self._sanitize_sheet_name(translate_fn(rs.name))
            else:
                ws.title = rs.name

            # Dịch cell values
            for row_idx in range(rs.nrows):
                for col_idx in range(rs.ncols):
                    value = rs.cell_value(row_idx, col_idx)
                    if isinstance(value, str) and value.strip():
                        ws.cell(row=row_idx + 1, column=col_idx + 1, value=translate_fn(value))
                    elif value != "":
                        ws.cell(row=row_idx + 1, column=col_idx + 1, value=value)

        actual_output = output_path.with_suffix(".xlsx")
        wb.save(actual_output)
        wb.close()
        logger.info("Đã ghi file xls -> xlsx: %s", safe_file_label(actual_output))
